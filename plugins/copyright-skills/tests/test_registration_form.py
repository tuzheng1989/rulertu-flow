"""registration_form.py 行为基线（B3）。

锚点（registration_form.py）：
- :48-49  fill 输出已存在拒绝覆盖
- :50-51  输出与模板同路径拒绝
- :59-64  changes.json 逐格记录 before/after
- :24-44  inspect 输出含 merged_ranges 与非空单元格清单

注意：同路径分支（:50-51）只有在 output 不存在时才可达——output 已存在会先命中
覆盖检查（:48-49），且路径相同时文件必然已存在；因此用"模板与输出指向同一条
尚不存在的路径"触发该分支，锁定检查语序与文案。
"""

from __future__ import annotations

import json

from openpyxl import Workbook, load_workbook


def make_template(path, values=None, merged=None):
    workbook = Workbook()
    sheet = workbook.active
    for coordinate, value in (values or {}).items():
        sheet[coordinate] = value
    for rng in merged or []:
        sheet.merge_cells(rng)
    workbook.save(path)
    return path


def run_fill(run_cli, scripts, template, facts, output):
    return run_cli(
        scripts.registration_form, "fill",
        "--template", template, "--facts", facts, "--output", output,
    )


def test_fill_existing_output_rejected(tmp_path, run_cli, scripts):
    """输出文件已存在：拒绝覆盖（锚点 :48-49，全局红线 1）。"""
    template = make_template(tmp_path / "template.xlsx", {"A1": "原值"})
    facts = tmp_path / "facts.json"
    facts.write_text(json.dumps({"values": {"A1": "新值"}}), encoding="utf-8")
    output = tmp_path / "out.xlsx"
    output.write_bytes(b"existing")

    result = run_fill(run_cli, scripts, template, facts, output)
    assert result.returncode != 0
    assert "refusing to overwrite output" in result.stderr
    assert output.read_bytes() == b"existing"


def test_same_template_and_output_path_rejected(tmp_path, run_cli, scripts):
    """模板与输出同路径：拒绝（锚点 :50-51，见文件头注释关于触发条件说明）。"""
    same = tmp_path / "target.xlsx"
    facts = tmp_path / "facts.json"
    facts.write_text(json.dumps({"values": {}}), encoding="utf-8")

    result = run_fill(run_cli, scripts, same, facts, same)
    assert result.returncode != 0
    assert "output must differ from template" in result.stderr
    assert not same.exists()  # 在写任何文件之前就拒绝


def test_fill_records_before_and_after(tmp_path, run_cli, scripts):
    """changes.json 逐格记录 before/after（锚点 :59-64, :73-75）。"""
    template = make_template(tmp_path / "template.xlsx", {"A1": "旧名称"})
    facts = tmp_path / "facts.json"
    facts.write_text(
        json.dumps({"values": {"A1": "新名称", "C3": 7}}, ensure_ascii=False),
        encoding="utf-8",
    )
    output = tmp_path / "out.xlsx"

    result = run_fill(run_cli, scripts, template, facts, output)
    assert result.returncode == 0, result.stderr

    audit = json.loads(output.with_suffix(output.suffix + ".changes.json").read_text(encoding="utf-8"))
    assert audit["sheet"] == "Sheet"
    # 独立真值：模板里 A1="旧名称"、C3 为空
    assert {"cell": "A1", "before": "旧名称", "after": "新名称"} in audit["changes"]
    assert {"cell": "C3", "before": None, "after": 7} in audit["changes"]

    filled = load_workbook(output)
    assert filled["Sheet"]["A1"].value == "新名称"
    assert filled["Sheet"]["C3"].value == 7
    # 模板文件本身未被改动
    assert load_workbook(template)["Sheet"]["A1"].value == "旧名称"


def test_inspect_reports_merged_ranges_and_nonempty_cells(tmp_path, run_cli, scripts):
    """inspect 输出含 merged_ranges 与非空单元格清单（锚点 :28-41）。"""
    template = make_template(
        tmp_path / "template.xlsx",
        {"A1": "软件全称", "B4": 42},
        merged=["A1:B3"],
    )
    output = tmp_path / "inspect.json"

    result = run_cli(
        scripts.registration_form, "inspect",
        "--template", template, "--output", output,
    )
    assert result.returncode == 0, result.stderr

    payload = json.loads(output.read_text(encoding="utf-8"))
    sheet = payload["sheets"][0]
    assert sheet["title"] == "Sheet"
    assert sheet["merged_ranges"] == ["A1:B3"]
    cells = {item["cell"]: item["value"] for item in sheet["nonempty_cells"]}
    # 独立真值：只写入了 A1 与 B4 两个非空格
    assert cells == {"A1": "软件全称", "B4": 42}
    assert sheet["image_count"] == 0
