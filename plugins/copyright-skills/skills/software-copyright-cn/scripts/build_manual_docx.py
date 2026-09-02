#!/usr/bin/env python3
"""Build an illustrated Chinese software manual from a page-oriented JSON spec."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError as exc:
    raise SystemExit("python-docx is required: python -m pip install python-docx") from exc


EXAMPLE = {
    "title": "软件名称 操作手册",
    "version": "V1.0",
    "date": "2026年9月",
    "pages": [
        {
            "title": "1 系统入口",
            "paragraphs": ["启动服务后，在浏览器中打开系统地址。", "页面显示主工作区。"],
            "steps": ["确认服务状态。", "打开系统首页。"],
            "image": "screenshots/01-home.png",
            "caption": "图1 系统首页",
            "evidence": ["poc/server/app.py", "poc/server/static/index.html"]
        }
    ]
}


def set_font(run, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_text(paragraph, text: str, size: float = 11, bold: bool = False) -> None:
    set_font(paragraph.add_run(text), "宋体", size, bold)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--spec", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--min-pages", type=int, default=10)
    parser.add_argument("--example", action="store_true")
    args = parser.parse_args()
    if args.example:
        print(json.dumps(EXAMPLE, ensure_ascii=False, indent=2))
        return 0
    if not args.spec or not args.output:
        parser.error("--spec and --output are required unless --example is used")
    output = args.output.resolve()
    if output.exists():
        raise SystemExit(f"refusing to overwrite output: {output}")
    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    pages = spec.get("pages", [])
    total_pages = 2 + len(pages)
    if total_pages < args.min_pages:
        raise SystemExit(
            f"manual has only {total_pages} explicit pages (cover + contents + {len(pages)} content); "
            f"at least {args.min_pages} required"
        )
    image_count = sum(bool(page.get("image")) for page in pages)
    if image_count == 0:
        raise SystemExit("manual must contain real screenshots or clearly labelled diagrams")

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(150)
    add_text(paragraph, spec["title"], 24, True)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(paragraph, spec.get("version", ""), 16)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(180)
    add_text(paragraph, spec.get("date", ""), 12)
    document.add_page_break()

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(heading, "目录", 18, True)
    for index, page in enumerate(pages, start=3):
        paragraph = document.add_paragraph()
        add_text(paragraph, f"{page['title']}  ……  {index}", 11)
    document.add_page_break()

    base = args.spec.resolve().parent
    for index, page in enumerate(pages):
        heading = document.add_paragraph()
        add_text(heading, page["title"], 16, True)
        for text in page.get("paragraphs", []):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(0.74)
            paragraph.paragraph_format.line_spacing = 1.5
            add_text(paragraph, text)
        for step_number, step in enumerate(page.get("steps", []), start=1):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.line_spacing = 1.35
            add_text(paragraph, f"{step_number}. {step}")
        image_value = page.get("image")
        if image_value:
            image_path = Path(image_value)
            if not image_path.is_absolute():
                image_path = base / image_path
            if not image_path.is_file():
                raise SystemExit(f"image not found: {image_path}")
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Cm(15.5))
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(caption, page.get("caption", image_path.name), 9)
        if page.get("evidence"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(8)
            add_text(paragraph, "材料复核依据：" + "；".join(page["evidence"]), 8)
        if index < len(pages) - 1:
            document.add_page_break()

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    audit = {
        "spec": str(args.spec.resolve()),
        "explicit_pages": total_pages,
        "content_pages": len(pages),
        "illustrated_content_pages": image_count,
        "evidence_references": sum(len(page.get("evidence", [])) for page in pages),
    }
    output.with_suffix(output.suffix + ".audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"explicit_pages={total_pages} illustrated_pages={image_count} output={output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

