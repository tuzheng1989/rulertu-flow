#!/usr/bin/env python3
"""Create an auditable Word source-code excerpt from an inventory JSON.

申报口径：源代码文档按原始物理行逐字渲染，保留注释与空行（中文功能注释是
审查要点）；行数达标按代码行统计——代码行 = 非注释、非空行的物理行，由
strip_source.code_line_rows 计算。3,000 行阈值判断、前/后 1,500 段截取、
登记表"源程序量"一律以代码行口径为准；manifest 双轨记录物理行数
（total_physical_lines）与代码行数（total_code_lines）及每文件统计
（strip_stats）。左列显示 相对路径:原物理行号。
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError as exc:
    raise SystemExit("python-docx is required: python -m pip install python-docx") from exc

from strip_source import StripError, code_line_rows


def decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def set_cell_shading(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--inventory", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--title", required=True)
    parser.add_argument("--version", required=True)
    parser.add_argument("--lines-per-page", type=int, default=50)
    parser.add_argument(
        "--first",
        action="append",
        metavar="PATH",
        help="把清单内该相对路径的文件前置到串接序列最前（可重复，按传入顺序排列）",
    )
    args = parser.parse_args()
    if args.lines_per_page < 1:
        raise SystemExit("--lines-per-page must be positive")
    output = args.output.resolve()
    if output.exists():
        raise SystemExit(f"refusing to overwrite output: {output}")

    inventory = json.loads(args.inventory.read_text(encoding="utf-8"))
    repo = Path(inventory["repo"])
    included = inventory["included_files"]
    # --first：按传入顺序把命中文件移到串接序列最前，其余保持清单相对顺序
    first_paths = list(dict.fromkeys(args.first or []))
    known = {entry["path"] for entry in included}
    unknown = [value for value in first_paths if value not in known]
    if unknown:
        available = ", ".join(entry["path"] for entry in included)
        raise SystemExit(
            "--first path not in inventory: " + ", ".join(unknown)
            + "; available paths: " + available
        )
    ordered = [entry for path in first_paths for entry in included if entry["path"] == path]
    ordered += [entry for entry in included if entry["path"] not in first_paths]

    records: list[dict] = []
    strip_stats: dict[str, dict] = {}
    total_original = 0
    total_code = 0
    for entry in ordered:
        path = repo / entry["path"]
        data = path.read_bytes()
        digest = hashlib.sha256(data).hexdigest()
        if digest != entry["sha256"]:
            raise SystemExit(f"source changed after inventory: {entry['path']}")
        text = decode(data)
        try:
            rows, file_stats = code_line_rows(text, path.suffix)
        except StripError as exc:
            raise SystemExit(f"strip failed: {entry['path']}: {exc}") from exc
        strip_stats[entry["path"]] = file_stats
        total_original += file_stats["original_lines"]
        total_code += len(rows)
        code_set = set(rows)
        # 渲染按原始物理行逐字保留（含注释与空行）；code 标记供代码行截断用
        for number, content in enumerate(text.splitlines(), start=1):
            records.append({
                "path": entry["path"],
                "line": number,
                "text": content,
                "code": number in code_set,
            })

    total = total_code
    if total <= 3000:
        selected = records
        head_cutoff = tail_start_index = None
    else:
        # 前 1,500 个代码行的物理行段 + 后 1,500 个代码行的物理行段，
        # 段内注释与空行随物理行一并保留；两段之间的物理行不提交
        head_cutoff = tail_start_index = None
        code_seen = 0
        for index, record in enumerate(records):
            if record["code"]:
                code_seen += 1
                if code_seen == 1500:
                    head_cutoff = index
                if code_seen == total - 1499:
                    tail_start_index = index
                    break
        selected = records[:head_cutoff + 1] + records[tail_start_index:]
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.1)
    section.right_margin = Cm(1.1)
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0.4)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run(f"{args.title} {args.version}  源代码")
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")

    pages = [selected[i:i + args.lines_per_page] for i in range(0, len(selected), args.lines_per_page)]
    for page_index, page in enumerate(pages):
        table = document.add_table(rows=0, cols=2)
        table.autofit = False
        table.columns[0].width = Cm(5.0)
        table.columns[1].width = Cm(13.5)
        for item in page:
            cells = table.add_row().cells
            row = cells[0]._tc.getparent()
            tr_props = row.get_or_add_trPr()
            height = OxmlElement("w:trHeight")
            height.set(qn("w:val"), "235")
            height.set(qn("w:hRule"), "atLeast")
            tr_props.append(height)
            cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cells[0], "F2F2F2")
            for cell, value, font_name, size in (
                (cells[0], f"{item['path']}:{item['line']}", "宋体", 6.5),
                (cells[1], item["text"].replace("\t", "    "), "Consolas", 6.5),
            ):
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                run = paragraph.add_run(value)
                run.font.name = font_name
                run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
                run.font.size = Pt(size)
        if page_index < len(pages) - 1:
            document.add_page_break()

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)

    def public(record: dict | None) -> dict | None:
        """manifest 边界条目不携带 code 标记，保持 {path, line, text} 三键。"""
        if record is None:
            return None
        return {"path": record["path"], "line": record["line"], "text": record["text"]}

    if total <= 3000:
        head_end = selected[min(1499, len(selected) - 1)] if selected else None
        tail_start = None
    else:
        head_end = records[head_cutoff]
        tail_start = records[tail_start_index]
    selected_code = sum(1 for record in selected if record["code"])
    manifest = {
        "inventory": str(args.inventory.resolve()),
        "title": args.title,
        "version": args.version,
        "total_physical_lines": total_original,
        "total_code_lines": total,
        "strip_stats": strip_stats,
        "selected_physical_line_count": len(selected),
        "selected_code_line_count": selected_code,
        "selection": "all" if total <= 3000 else "first-1500-code-lines-and-last-1500",
        "reordered_first": first_paths,
        "first": public(selected[0]) if selected else None,
        "head_end": public(head_end),
        "tail_start": public(tail_start),
        "last": public(selected[-1]) if selected else None,
        "selected_sha256": hashlib.sha256(
            json.dumps(
                [public(record) for record in selected],
                ensure_ascii=False, separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest(),
        "pages_at_configured_line_count": len(pages),
        "lines_per_page": args.lines_per_page,
    }
    output.with_suffix(output.suffix + ".manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        f"physical_lines={total_original} code_lines={total} "
        f"selected_physical={len(selected)} selected_code={selected_code} "
        f"pages={len(pages)} output={output}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

