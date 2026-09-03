#!/usr/bin/env python3
"""Inspect or safely fill a copied Excel worksheet for the online registration form."""

from __future__ import annotations

import argparse
import json
import shutil
from datetime import date, datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("openpyxl is required: python -m pip install openpyxl") from exc


def serializable(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def inspect(template: Path, output: Path) -> None:
    workbook = load_workbook(template, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        cells = [
            {"cell": cell.coordinate, "value": serializable(cell.value)}
            for row in sheet.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        ]
        sheets.append({
            "title": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "merged_ranges": [str(value) for value in sheet.merged_cells.ranges],
            "nonempty_cells": cells,
            "image_count": len(getattr(sheet, "_images", [])),
        })
    payload = {"template": str(template.resolve()), "sheets": sheets}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def fill(template: Path, facts_path: Path, output: Path) -> None:
    if output.exists():
        raise SystemExit(f"refusing to overwrite output: {output}")
    if template.resolve() == output.resolve():
        raise SystemExit("output must differ from template")
    facts = json.loads(facts_path.read_text(encoding="utf-8"))
    shutil.copy2(template, output)
    workbook = load_workbook(output, data_only=False)
    sheet_name = facts.get("sheet") or workbook.sheetnames[0]
    if sheet_name not in workbook.sheetnames:
        raise SystemExit(f"sheet not found: {sheet_name}")
    sheet = workbook[sheet_name]
    changes = []
    for coordinate, value in facts.get("values", {}).items():
        cell = sheet[coordinate]
        before = serializable(cell.value)
        cell.value = value
        changes.append({"cell": coordinate, "before": before, "after": serializable(value)})
    workbook.save(output)
    audit = {
        "template": str(template.resolve()),
        "output": str(output.resolve()),
        "sheet": sheet_name,
        "changes": changes,
        "warning": "Open the output in Excel or LibreOffice to verify drawings, controls, print area, and pagination.",
    }
    output.with_suffix(output.suffix + ".changes.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)
    inspect_parser = subparsers.add_parser("inspect")
    inspect_parser.add_argument("--template", type=Path, required=True)
    inspect_parser.add_argument("--output", type=Path, required=True)
    fill_parser = subparsers.add_parser("fill")
    fill_parser.add_argument("--template", type=Path, required=True)
    fill_parser.add_argument("--facts", type=Path, required=True)
    fill_parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.command == "inspect":
        inspect(args.template.resolve(), args.output.resolve())
    else:
        args.output.resolve().parent.mkdir(parents=True, exist_ok=True)
        fill(args.template.resolve(), args.facts.resolve(), args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
