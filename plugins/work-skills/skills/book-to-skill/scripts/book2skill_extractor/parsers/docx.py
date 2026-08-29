from __future__ import annotations

import zipfile
from pathlib import Path
from xml.etree import ElementTree


def _validate_xml(archive: zipfile.ZipFile) -> None:
    for name in archive.namelist():
        if name.lower().endswith((".xml", ".rels")):
            raw = archive.read(name).upper()
            if b"<!DOCTYPE" in raw or b"<!ENTITY" in raw:
                raise RuntimeError(f"Unsafe XML declaration in DOCX member: {name}")


def _stdlib_docx(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        _validate_xml(archive)
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    parts: list[str] = []
    body = root.find(f"{ns}body")
    for block in list(body if body is not None else root):
        if block.tag == f"{ns}p":
            text = "".join(node.text or "" for node in block.iter(f"{ns}t"))
            if text:
                parts.append(text)
        elif block.tag == f"{ns}tbl":
            for row in block.iter(f"{ns}tr"):
                cells = [
                    "".join(node.text or "" for node in cell.iter(f"{ns}t")).strip()
                    for cell in row.iter(f"{ns}tc")
                ]
                if any(cells):
                    parts.append("\t".join(cells))
    return "\n".join(parts)


def extract_docx(path: Path) -> tuple[str, str]:
    with zipfile.ZipFile(path) as archive:
        _validate_xml(archive)
    try:
        import docx
    except ImportError:
        return _stdlib_docx(path), "zipfile-docx"
    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts), "python-docx"
